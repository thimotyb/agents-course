#!/usr/bin/env python3
"""Generate the client-facing course syllabus from the canonical local sources."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import html


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "AI-Agent-Programming-Course-Syllabus.docx"
BLUE = "123B5D"
TEAL = "00A6A6"
PALE = "EAF4F7"
INK = "183042"
MUTED = "5C6F7D"

CANONICAL_TITLES = {
    "M1": "Getting Started with LLMs and LangChain",
    "M2": "Summarization with LangChain and LangGraph",
    "M3": "RAG Systems",
    "M4": "Advanced RAG",
    "M5": "AI Agents Architectures with LangGraph",
    "M6": "Google Agent Development Kit",
    "M7": "Labs and Exercises from Building LLM Applications",
}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field(paragraph, field: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def text_of(node) -> str:
    return " ".join(node.text_content().split())


def section_payload(doc_tree, heading: str) -> list[str]:
    matches = doc_tree.xpath(
        "//*[self::h2 or self::h3][normalize-space(.)=$heading]", heading=heading
    )
    if not matches:
        return []
    current = matches[0].getnext()
    values: list[str] = []
    while current is not None and current.tag not in {"h2", "h3"}:
        if current.tag in {"p", "ul", "ol"}:
            if current.tag in {"ul", "ol"}:
                values.extend(text_of(li) for li in current.xpath("./li") if text_of(li))
            elif text_of(current):
                values.append(text_of(current))
        current = current.getnext()
    return values


def extract_modules() -> list[dict]:
    modules = []
    for path in sorted((ROOT / "site/chapters").glob("chapter-*.html")):
        tree = html.fromstring(path.read_text(encoding="utf-8"))
        title = text_of(tree.xpath("//h1")[0])
        hero = text_of(tree.xpath("//*[contains(concat(' ', normalize-space(@class), ' '), ' hero-copy ')]")[0])
        numbered = []
        for heading in tree.xpath("//article//*[self::h2 or self::h3]"):
            value = text_of(heading)
            if re.match(r"^\d+(?:\.\d+)?\.?(?:\s|$)", value):
                numbered.append(value)
        code = title.split(" - ", 1)[0]
        modules.append(
            {
                "code": code,
                "title": CANONICAL_TITLES[code],
                "hero": hero,
                "outcomes": section_payload(tree, "Learning outcomes"),
                "topics_summary": section_payload(tree, "Topics"),
                "topics": numbered,
                "lab": section_payload(tree, "Execution steps"),
                "practice": section_payload(tree, "Practice tasks"),
                "deliverable": section_payload(tree, "Submission output"),
                "takeaways": section_payload(tree, "Key takeaways"),
            }
        )
    return modules


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color, before, after in (
        ("Title", 30, BLUE, 0, 12),
        ("Subtitle", 14, MUTED, 0, 18),
        ("Heading 1", 22, BLUE, 14, 8),
        ("Heading 2", 15, TEAL, 12, 5),
        ("Heading 3", 11, BLUE, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Aptos Display" if name != "Heading 3" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Subtitle"
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Module Label" not in styles:
        style = styles.add_style("Module Label", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style.font.size = Pt(10)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(TEAL)
        style.paragraph_format.space_after = Pt(4)


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.1)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    header = section.header.paragraphs[0]
    header.text = "AI AGENT PROGRAMMING  |  COURSE SYLLABUS"
    header.style = document.styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor.from_string(MUTED)
    header.runs[0].font.size = Pt(8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("Client presentation  •  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_field(footer, "PAGE")


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.add_run(item)


def add_label(document: Document, label: str) -> None:
    document.add_paragraph(label.upper(), style="Module Label")


def build_document(modules: list[dict]) -> Document:
    document = Document()
    configure_styles(document)
    configure_page(document)
    document.core_properties.title = "AI Agent Programming Course Syllabus"
    document.core_properties.subject = "Client-facing syllabus for the AI Agent Programming course"
    document.core_properties.author = "Course Team"
    document.core_properties.keywords = "AI agents, LangChain, LangGraph, RAG, Google ADK, Ollama"

    # Cover
    document.add_paragraph("CLIENT COURSE SYLLABUS", style="Module Label")
    document.add_paragraph("AI Agent Programming", style="Title")
    document.add_paragraph("LangChain • LangGraph • RAG • Google ADK", style="Subtitle")
    rule = document.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.LEFT
    rule.autofit = False
    rule.columns[0].width = Inches(6.2)
    cell = rule.cell(0, 0)
    set_cell_shading(cell, TEAL)
    cell.height = Cm(0.16)
    cell.text = ""
    document.add_paragraph()
    lead = document.add_paragraph()
    lead.paragraph_format.space_before = Pt(18)
    lead.paragraph_format.space_after = Pt(18)
    lead.add_run(
        "A practical, module-based programme for designing, implementing, "
        "evaluating, and operating LLM applications and agentic systems."
    ).font.size = Pt(15)

    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    labels = ("Audience", "Primary language", "Delivery format", "Document date")
    values = (
        "Developers and technical architects with Python basics",
        "English (Italian available on the interactive course site)",
        "Instructor-led learning, guided labs, exercises, and practical deliverables",
        date.today().strftime("%d %B %Y"),
    )
    for row, label, value in zip(meta.rows, labels, values):
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(12.0)
        set_cell_shading(row.cells[0], PALE)
        for cell in row.cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(value)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Course profile
    document.add_heading("Course profile", level=1)
    document.add_heading("Purpose", level=2)
    document.add_paragraph(
        "The course develops the capabilities required to move from foundational LLM applications "
        "to retrieval-augmented generation, agent orchestration, tool integration, memory, "
        "guardrails, observability, and production-oriented delivery."
    )
    document.add_heading("Overall learning outcomes", level=2)
    add_bullets(
        document,
        [
            "Design LLM applications using LangChain and graph-based workflows using LangGraph.",
            "Implement foundational and advanced RAG pipelines for knowledge-grounded applications.",
            "Build tool-using agents and reason about state, routing, memory, and control flow.",
            "Apply reliability practices including tracing, error handling, guardrails, and evaluation.",
            "Create runnable, documented implementations through guided labs and practical assignments.",
        ],
    )
    document.add_heading("Technical environment", level=2)
    add_bullets(
        document,
        [
            "Ubuntu development environment with Python basics.",
            "Local Ollama runtime for classroom examples.",
            "Default local model: gemma4:e4b.",
            "No OpenAI API key is required for the baseline classroom examples.",
        ],
    )
    document.add_heading("Learning approach", level=2)
    add_bullets(
        document,
        [
            "Conceptual framing followed by implementation-oriented examples.",
            "Guided labs that connect architecture, code, execution, and observation.",
            "Exercises focused on adaptation, debugging, and measurable outputs.",
            "Reproducible deliverables with setup instructions, commands, configuration, and results.",
        ],
    )

    document.add_heading("Programme at a glance", level=1)
    table = document.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("Module", "Title", "Focus")
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, BLUE)
        run = cell.paragraphs[0].add_run(value)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_margins(cell)
    set_repeat_table_header(table.rows[0])
    for module in modules:
        cells = table.add_row().cells
        focus = module["hero"]
        for cell, value in zip(cells, (module["code"], module["title"], focus)):
            cell.text = value
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

    # Module sheets
    for module in modules:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph(module["code"], style="Module Label")
        document.add_heading(module["title"], level=1)
        intro = document.add_paragraph()
        intro.paragraph_format.space_after = Pt(12)
        intro.add_run(module["hero"]).bold = True

        add_label(document, "Learning outcomes")
        add_bullets(document, module["outcomes"])

        add_label(document, "Syllabus")
        if module["topics"]:
            for topic in module["topics"]:
                if re.match(r"^\d+\.\d+", topic):
                    paragraph = document.add_paragraph(style="List Bullet 2")
                else:
                    paragraph = document.add_paragraph()
                    paragraph.paragraph_format.space_before = Pt(5)
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run(topic).bold = True
                    continue
                paragraph.paragraph_format.space_after = Pt(1)
                paragraph.add_run(topic)
        else:
            add_bullets(document, module["topics_summary"])

        if module["lab"]:
            add_label(document, "Guided lab")
            document.add_paragraph(module["lab"][0])
        if module["practice"]:
            add_label(document, "Practice")
            document.add_paragraph(module["practice"][0])
        if module["deliverable"]:
            add_label(document, "Expected deliverable")
            document.add_paragraph(module["deliverable"][0])

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Completion profile", level=1)
    document.add_paragraph(
        "On completion, participants will have progressed from foundational LLM application patterns "
        "to the design and implementation of retrieval systems and production-oriented agents. "
        "The programme closes with reproducible practical work that demonstrates architecture, "
        "implementation, execution, debugging, and documentation skills."
    )
    document.add_heading("Course outputs", level=2)
    add_bullets(
        document,
        [
            "Runnable examples and guided lab results.",
            "RAG and agent workflow implementations.",
            "An ADK-based agent with operational documentation.",
            "A reproducible final lab report describing commands, configuration, outputs, and trade-offs.",
        ],
    )
    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    run = note.add_run(
        "Scope note: delivery duration, scheduling, assessment weighting, and commercial terms are "
        "intentionally excluded and can be defined in the training proposal."
    )
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(MUTED)
    return document


def main() -> None:
    modules = extract_modules()
    if len(modules) != 7:
        raise SystemExit(f"Expected 7 modules, found {len(modules)}")
    document = build_document(modules)
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()
